"""
Created as part of the C-NEMS Project.

Written by:  Sauleh Siddiqui
Contact:  sauleh@american.edu
Created on:  10/2/26

Builds the C-CDM input tables in ``input/cdm`` from their published sources.

C-CDM does the job of the Commercial Demand Module (CDM) in EIA's National Energy Modeling System
(NEMS, https://github.com/EIAgov/NEMS, release tag ``AEO2026-Public-Release``, Apache License
2.0). This script reads the NEMS CDM input files C-CDM needs and rows of EIA's Annual Energy
Outlook 2026 (AEO2026) result tables, and writes them as tidy CSV files whose ``#`` header lines
name the source file, the lines or series codes read, the units and any transformation. The NEMS
text files are not shipped as they are: the repository normalizes line endings and trailing
whitespace, which would rewrite them.

The model does not run this script. Run it from the repository root, with ``openpyxl``::

    pixi exec --spec python=3.12 --spec openpyxl -- python src/models/cdm/build_inputs.py
        --reference <AEO2026 reference workbook> --high-demand <high electricity demand workbook>
        --supplement <AEO2026 supplemental residential/commercial/industrial workbook>

The two year-by-year workbooks hold AEO2026 Tables 3 and 5 for the Counterfactual Baseline case
(NEMS run ``cb2026.d021826b``) and the High Electricity Demand case (``higheldmd.d021526a``); the
supplemental workbook holds Table 22 for the Counterfactual Baseline. EIA publishes them on its
AEO2026 tables page, https://www.eia.gov/outlooks/aeo/tables_ref.php: the year-by-year workbook of
the reference case (``yearbyyear.xlsx``), the supplemental residential, commercial and industrial
workbook (``sup_rci.xlsx``) and, through the side-case tables it links to, the High Electricity
Demand case. Rows are found by their EIA series code, never by position. NEMS files are downloaded
from the release tag unless ``--nems-dir`` points at a local copy of ``input/cdm``.
"""

import argparse
import csv
import re
import urllib.request
from pathlib import Path

NEMS_TAG = 'AEO2026-Public-Release'
NEMS_URL = f'https://raw.githubusercontent.com/EIAgov/NEMS/{NEMS_TAG}/input/cdm'
YEARS = range(2025, 2051)

REGIONS = [
    'new_england', 'middle_atlantic', 'east_north_central', 'west_north_central',
    'south_atlantic', 'east_south_central', 'west_south_central', 'mountain', 'pacific',
]  # fmt: skip
BUILDINGS = [
    'assembly', 'education', 'food_sales', 'food_service', 'health_care', 'lodging',
    'large_office', 'small_office', 'mercantile_service', 'warehouse', 'other',
]  # fmt: skip
SERVICES = [
    'heating', 'cooling', 'water_heating', 'ventilation', 'cooking', 'lighting',
    'refrigeration', 'data_center_servers', 'computers_office_equipment', 'other_uses',
]  # fmt: skip
FUELS = {'EL': 'electricity', 'NG': 'natural_gas', 'DS': 'distillate'}
# kmels.txt [MelsElQ] order, which is also MelsELQ(1..20) in comm.f:3488-3528
MELS = [
    'distribution_transformers', 'kitchen_ventilation', 'security_systems',
    'lab_refrigerators_freezers', 'medical_imaging', 'large_video_boards', 'coffee_brewers',
    'off_road_electric_vehicles', 'fume_hoods', 'laundry', 'elevators', 'escalators',
    'it_equipment', 'uninterruptible_power_supplies', 'shredders', 'private_branch_exchange',
    'voice_over_ip', 'point_of_sale', 'warehouse_robots', 'televisions',
]  # fmt: skip
CASES = {'reference': 'kmels.txt', 'high_electricity_demand': 'kmels_higheldmd.txt'}

# AEO2026 Table 5 series codes (CKI000), by fuel and service
TABLE5 = {
    ('electricity', 'heating'): 'CKI000:ga_SpaceHeating',
    ('electricity', 'cooling'): 'CKI000:ga_SpaceCooling',
    ('electricity', 'water_heating'): 'CKI000:ga_WaterHeating',
    ('electricity', 'ventilation'): 'CKI000:ga_Ventilation',
    ('electricity', 'cooking'): 'CKI000:ga_Cooking',
    ('electricity', 'lighting'): 'CKI000:ga_Lighting',
    ('electricity', 'refrigeration'): 'CKI000:ga_Refrigeration',
    ('electricity', 'data_center_servers'): 'CKI000:ga_DataCtrServ',
    ('electricity', 'computers_office_equipment'): 'CKI000:ha_CompOfficeEquipme',
    ('electricity', 'other_uses'): 'CKI000:ha_OtherUses',
    ('electricity', 'gross_end_use_total'): 'CKI000:ha_ElecSubtotal',
    ('electricity', 'ev_charging'): 'CKI000:ha_PurchElecEVCha',
    ('electricity', 'own_generation'): 'CKI000:ha_OwnGeneration',
    ('electricity', 'purchased_total'): 'CKI000:ha_PurchasedElec',
    ('natural_gas', 'heating'): 'CKI000:ia_SpaceHeating',
    ('natural_gas', 'cooling'): 'CKI000:ia_SpaceCooling',
    ('natural_gas', 'water_heating'): 'CKI000:ia_WaterHeating',
    ('natural_gas', 'cooking'): 'CKI000:ia_Cooking',
    ('natural_gas', 'other_uses'): 'CKI000:ia_OtherUses',
    ('natural_gas', 'delivered_total'): 'CKI000:ia_DeliveredEner',
}
TABLE22_FLOORSPACE = {
    'assembly': 'CST000:ca_Assembly',
    'education': 'CST000:ca_Education',
    'food_sales': 'CST000:ca_FoodSales',
    'food_service': 'CST000:ca_FoodService',
    'health_care': 'CST000:ca_HealthCare',
    'lodging': 'CST000:ca_Lodging',
    'large_office': 'CST000:ca_Office-Large',
    'small_office': 'CST000:ca_Office-Small',
    'mercantile_service': 'CST000:ca_Mercantile/Se',
    'warehouse': 'CST000:ca_Warehouse',
    'other': 'CST000:ca_Other',
}
TABLE22_CHP_GAS = 'CST000:dge_NaturalGas'
AEO_PAGE = 'https://www.eia.gov/outlooks/aeo/tables_ref.php'
TABLE3_PRICES = {'natural_gas': 'PRC000:ca_NaturalGas', 'electricity': 'PRC000:ca_Electricity'}

NUMBER = re.compile(r'-?\d+(?:\.\d*)?')


def _numbers(line: str) -> list[float]:
    return [float(x) for x in NUMBER.findall(line)]


def _nems_lines(name: str, nems_dir: Path | None) -> list[str]:
    """Read one NEMS CDM input file as lines, from ``nems_dir`` or the release tag."""
    if nems_dir is not None:
        raw = (nems_dir / name).read_bytes()
    else:
        with urllib.request.urlopen(f'{NEMS_URL}/{name}', timeout=60) as response:
            raw = response.read()
    return raw.decode('latin-1').splitlines()


def _tag_line(lines: list[str], tag: str) -> int:
    """Return the 0-based index of the one line holding a kmels ``[tag]`` header."""
    hits = [i for i, line in enumerate(lines) if f'[{tag}]' in line]
    if len(hits) != 1:
        raise ValueError(f'expected one [{tag}] header, found {len(hits)}')
    return hits[0]


def _write(path: Path, header: list[str], columns: list[str], rows: list[list]) -> None:
    """Write a tidy CSV with ``#`` comment lines above the column names."""
    with open(path, 'w', newline='', encoding='utf-8') as fh:
        fh.writelines(f'# {line}\n' if line else '#\n' for line in header)
        writer = csv.writer(fh, lineterminator='\n')
        writer.writerow(columns)
        writer.writerows(rows)
    print(f'{path.name}: {len(rows)} rows')


def build_floorspace(nems_dir: Path | None, out: Path) -> None:
    """Write 2018 floorspace by division and building type from ``kflspc.txt``."""
    lines = _nems_lines('kflspc.txt', nems_dir)
    found = [
        (i + 1, r)
        for i, r in enumerate(_numbers(line) for line in lines)
        if i >= 100 and len(r) == 9
    ]
    found = found[:110]
    rows = [r for _, r in found]
    if len(rows) != 110:
        raise ValueError(f'kflspc.txt: expected 110 cohort rows, found {len(rows)}')
    table: list[list] = []
    total = 0.0
    for b, building in enumerate(BUILDINGS):
        block = rows[b * 10 : (b + 1) * 10]
        for r, region in enumerate(REGIONS):
            value = round(sum(row[r] for row in block), 6)
            table.append([region, building, value])
            total += value
    _write(
        out / 'nems_floorspace_2018.csv',
        [
            'Commercial floorspace in 2018 by census division and building type, million sqft.',
            f'Source: NEMS input/cdm/kflspc.txt, release tag {NEMS_TAG} (EIA, CBECS 2018).',
            f'Read: numeric lines {found[0][0]}-{found[-1][0]}, one block of 10 vintage-cohort',
            'rows per building type (file header order), one column per census division; summed',
            'over the 10 cohorts.',
            f'Total {total:,.1f} million sqft.',
        ],
        ['region', 'building_type', 'floorspace_million_sqft'],
        table,
    )


def build_eui(nems_dir: Path | None, out: Path) -> None:
    """Write 2018 energy use intensity by division, building, service and fuel from kintens."""
    pattern = re.compile(r'\s*((?:-?\d+(?:\.\d*)?\s+){10})(EL|NG|DS)\b')
    recs = []
    for line in _nems_lines('kintens.txt', nems_dir)[100:]:
        m = pattern.match(line)
        if m:
            recs.append((_numbers(m.group(1)), m.group(2)))
    if len(recs) != 9 * 11 * 3:
        raise ValueError(f'kintens.txt: expected 297 fuel lines, found {len(recs)}')
    table, i = [], 0
    for region in REGIONS:
        for building in BUILDINGS:
            for code in ('EL', 'NG', 'DS'):
                values, fuel = recs[i]
                i += 1
                if fuel != code:
                    raise ValueError(f'kintens.txt: fuel order broken at record {i}')
                for s, service in enumerate(SERVICES):
                    table.append([region, building, service, FUELS[code], values[s]])
    _write(
        out / 'nems_eui_2018.csv',
        [
            'Energy use intensity in 2018 (CBECS 2018), thousand Btu per sqft.',
            f'Source: NEMS input/cdm/kintens.txt, release tag {NEMS_TAG}.',
            'Read: data lines 101+, blocks by census division, then building type, then one line',
            'each for electricity (EL), natural gas (NG) and distillate (DS); columns are the 10',
            'services. Distillate is kept only because NEMS uses it in the other-uses fuel split.',
        ],
        ['region', 'building_type', 'service', 'fuel', 'eui_kbtu_per_sqft'],
        table,
    )


def _kmels_case(lines: list[str]) -> dict:
    """Read one kmels file: multipliers, explicit MEL share, server index, MEL base and index."""
    k = _tag_line(lines, 'dcf')
    first = next(i for i in range(k + 1, len(lines)) if len(_numbers(lines[i])) >= 10)
    dcf = [_numbers(lines[first + b])[:10] for b in range(11)]
    k_shr = _tag_line(lines, 'xplicitmiscshr')
    share = _numbers(lines[k_shr + 1])
    k_srv = _tag_line(lines, 'DataCtrServPenetration')
    srv_years = [int(y) for y in _numbers(lines[k_srv + 1])]
    srv = dict(zip(srv_years, _numbers(lines[k_srv + 2]), strict=True))
    k_q = _tag_line(lines, 'MelsElQ')
    base = [_numbers(lines[k_q + 1 + j])[0] for j in range(20)]
    k_i = _tag_line(lines, 'MarketPenetrationMels')
    years = [int(y) for y in _numbers(lines[k_i + 1])]
    index = [
        dict(zip(years, _numbers(lines[k_i + 2 + j])[: len(years)], strict=True)) for j in range(20)
    ]
    if len(share) != 11 or len(srv) != 32:
        raise ValueError('kmels: unexpected [xplicitmiscshr] or [DataCtrServPenetration] length')
    return {
        'dcf': dcf, 'dcf_line': first + 1, 'share': share, 'share_line': k_shr + 2,
        'srv': srv, 'srv_line': k_srv + 3, 'base': base, 'base_line': k_q + 2, 'index': index,
        'index_line': k_i + 3,
    }  # fmt: skip


def build_kmels(nems_dir: Path | None, out: Path) -> None:
    """Write the data-center multipliers, server index and explicit MEL inputs, by case."""
    dcf_rows, srv_rows, shr_rows, base_rows, idx_rows, notes = [], [], [], [], [], []
    for case, name in CASES.items():
        k = _kmels_case(_nems_lines(name, nems_dir))
        for b, building in enumerate(BUILDINGS):
            dcf_rows += [[case, building, s, k['dcf'][b][j]] for j, s in enumerate(SERVICES)]
            shr_rows.append([case, building, k['share'][b]])
        srv_rows += [[case, y, v] for y, v in sorted(k['srv'].items())]
        idx_rows += [
            [case, mel, y, v] for j, mel in enumerate(MELS) for y, v in k['index'][j].items()
        ]
        if case == 'reference':
            base_rows = [[mel, k['base'][j]] for j, mel in enumerate(MELS)]
        notes.append(
            f'{case}: {name} lines {k["dcf_line"]}-{k["dcf_line"] + 10} (multipliers), '
            f'{k["share_line"]} (explicit share), {k["srv_line"]} (server index), '
            f'{k["base_line"]}-{k["base_line"] + 19} (MEL base), {k["index_line"]}-'
            f'{k["index_line"] + 19} (MEL index)'
        )
    src = [f'Source: NEMS input/cdm/kmels.txt and kmels_higheldmd.txt, release tag {NEMS_TAG}.']
    _write(
        out / 'nems_dc_multipliers.csv',
        [
            'Data-center intensity multipliers [dcf] by building type and service, unitless.',
            *src,
            *notes,
            'Each service demand is scaled by (1 - d) + multiplier x d, d the data-center share',
            '(comm.f:3430-3434); where NEMS applies it is set in the model code, not here.',
        ],
        ['case', 'building_type', 'service', 'multiplier'],
        dcf_rows,
    )
    _write(
        out / 'nems_server_index.csv',
        [
            'Data-center server annual energy index [DataCtrServPenetration],',
            '2018 (CBECS year) = 1.',
            *src,
            *notes,
        ],
        ['case', 'year', 'server_index'],
        srv_rows,
    )
    _write(
        out / 'nems_mel_share.csv',
        [
            'Share of miscellaneous electric load explicitly modeled in the base year',
            '[xplicitmiscshr] by building type, unitless (comm.f:3729-3745 applies 1 - share).',
            *src,
            *notes,
        ],
        ['case', 'building_type', 'explicit_share'],
        shr_rows,
    )
    _write(
        out / 'nems_mel_base.csv',
        [
            'Base-year annual energy consumption of explicit miscellaneous electric loads',
            '[MelsElQ], Btu per sqft of eligible floorspace. The kmels header says trillion',
            'Btu/sqft/yr; the scaling in comm.f:3488-3528 (value / 1000 x billion sqft = trillion',
            'Btu) shows Btu/sqft.',
            *src,
            *notes,
            'Identical in both files; the reference file is read.',
        ],
        ['mel', 'base_btu_per_sqft'],
        base_rows,
    )
    _write(
        out / 'nems_mel_index.csv',
        [
            'Explicit miscellaneous electric load annual energy index [MarketPenetrationMels],',
            '2018 (CBECS year) = 1.',
            *src,
            *notes,
        ],
        ['case', 'mel', 'year', 'index'],
        idx_rows,
    )


def build_elasticity(nems_dir: Path | None, out: Path) -> None:
    """Write short-run price elasticities of service demand from ``ksdela.txt``."""
    lines = _nems_lines('ksdela.txt', nems_dir)
    found = [(i + 1, r) for i, r in enumerate(_numbers(line) for line in lines) if i >= 100 and r]
    rows = [r for _, r in found]
    if len(rows) != 90:
        raise ValueError(f'ksdela.txt: expected 90 rows, found {len(rows)}')
    table = []
    for r, region in enumerate(REGIONS):
        for s, service in enumerate(SERVICES):
            for f, fuel in enumerate(FUELS.values()):
                table.append([region, service, fuel, rows[r * 10 + s][f] + 0.0])
    _write(
        out / 'nems_price_elasticity.csv',
        [
            'Short-run price elasticity of service demand, unitless.',
            f'Source: NEMS input/cdm/ksdela.txt, release tag {NEMS_TAG}, numeric lines',
            f'{found[0][0]}-{found[-1][0]}: one block of 10 service rows per census division,',
            'columns electricity, natural gas, distillate. The file header still names services',
            '8 and 9 office PCs and non-PCs;',
            'NEMS reads rows by position, so rows 8 and 9 apply to servers and office equipment.',
        ],
        ['region', 'service', 'fuel', 'elasticity'],
        table,
    )


def _workbook_rows(path: Path, codes: set[str]) -> tuple[str, dict]:
    """Return the sheet name and {series code: (excel row, label, {year: value})} of a workbook."""
    import openpyxl  # pyrefly: ignore[missing-import]  - builder-only dependency, not in the env

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    year_col = {j: int(v) for j, v in enumerate(header) if isinstance(v, int | float)}
    found: dict[str, tuple] = {}
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        code = row[0]
        if code in codes:
            if code in found:
                raise ValueError(f'{path.name}: series {code} appears twice')
            label = next((str(v).strip() for v in row[1:] if isinstance(v, str)), '')
            values = {year_col[j]: row[j] for j in year_col if row[j] is not None}
            found[code] = (i, label, values)
    missing = codes - set(found)
    if missing:
        raise ValueError(f'{path.name}: series not found: {sorted(missing)}')
    return ws.title, found


def build_aeo(reference: Path, high: Path, supplement: Path, out: Path) -> None:
    """Write the AEO2026 floorspace, end-use consumption and price rows C-CDM reads."""
    sheet22, t22 = _workbook_rows(supplement, set(TABLE22_FLOORSPACE.values()) | {TABLE22_CHP_GAS})
    floor = []
    for building, code in TABLE22_FLOORSPACE.items():
        row, _, values = t22[code]
        floor += [[building, y, values[y], code, row] for y in YEARS]
    _write(
        out / 'aeo2026_floorspace.csv',
        [
            'Commercial floorspace by building type, billion sqft.',
            'Source: EIA, Annual Energy Outlook 2026, supplemental Table 22 (Commercial Sector',
            'Energy Consumption, Floorspace, ...), Counterfactual Baseline case, NEMS run',
            f'{sheet22}; series code and workbook row per line.',
            f'Downloaded from {AEO_PAGE}.',
            'Used for both cases (assumption: AEO2026 Table 5 total floorspace differs between',
            'the cases by under 0.01%).',
        ],
        ['building_type', 'year', 'floorspace_billion_sqft', 'series_code', 'source_row'],
        floor,
    )
    end_use, prices, sheets = [], [], []
    for case, path in (('reference', reference), ('high_electricity_demand', high)):
        codes = set(TABLE5.values()) | set(TABLE3_PRICES.values())
        sheet, rows = _workbook_rows(path, codes)
        sheets.append(f'{case}: NEMS run {sheet}')
        for (fuel, item), code in TABLE5.items():
            row, _, values = rows[code]
            end_use += [[case, fuel, item, y, values[y], code, f'Table 5 r{row}'] for y in YEARS]
        for fuel, code in TABLE3_PRICES.items():
            row, _, values = rows[code]
            prices += [[case, fuel, y, values[y], code, row] for y in YEARS]
    row, _, values = t22[TABLE22_CHP_GAS]
    for case in ('reference', 'high_electricity_demand'):
        end_use += [
            [case, 'natural_gas', 'combined_heat_power_input', y, values[y] / 1000.0,
             TABLE22_CHP_GAS, f'Table 22 r{row} (trillion Btu / 1000)']
            for y in YEARS
        ]  # fmt: skip
    _write(
        out / 'aeo2026_end_use.csv',
        [
            'Commercial consumption by fuel and end use, quadrillion Btu.',
            'Source: EIA, Annual Energy Outlook 2026, Table 5 (Commercial Sector Key Indicators',
            'and Consumption), by series code;',
            *sheets,
            f'Downloaded from {AEO_PAGE} (side cases through its side-case tables).',
            'Electricity rows are gross end use; gross_end_use_total, ev_charging, own_generation',
            'and purchased_total are the table rows that turn it into purchased electricity.',
            'combined_heat_power_input: natural gas input to commercial distributed generation and',
            'combined heat and power, supplemental Table 22, Counterfactual Baseline, trillion Btu',
            'divided by 1000; used for both cases (assumption). AEO2026 counts it in other uses.',
        ],
        ['case', 'fuel', 'service', 'year', 'consumption_quad', 'series_code', 'source_row'],
        end_use,
    )
    _write(
        out / 'aeo2026_prices.csv',
        [
            'Commercial delivered energy prices, 2025 USD per million Btu.',
            'Source: EIA, Annual Energy Outlook 2026, Table 3 (Energy Prices by Sector and',
            'Source), commercial rows, by series code;',
            *sheets,
            f'Downloaded from {AEO_PAGE} (side cases through its side-case tables).',
        ],
        ['case', 'fuel', 'year', 'price_2025usd_per_mmbtu', 'series_code', 'source_row'],
        prices,
    )


def main() -> None:
    """Parse the command line and write every C-CDM input table."""
    parser = argparse.ArgumentParser(description='Build the C-CDM input tables in input/cdm.')
    parser.add_argument('--reference', type=Path, required=True)
    parser.add_argument('--high-demand', type=Path, required=True)
    parser.add_argument('--supplement', type=Path, required=True)
    parser.add_argument('--nems-dir', type=Path, default=None)
    parser.add_argument('--out', type=Path, default=Path('input/cdm'))
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    build_floorspace(args.nems_dir, args.out)
    build_eui(args.nems_dir, args.out)
    build_kmels(args.nems_dir, args.out)
    build_elasticity(args.nems_dir, args.out)
    build_aeo(args.reference, args.high_demand, args.supplement, args.out)


if __name__ == '__main__':
    main()
